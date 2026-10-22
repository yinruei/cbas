#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
import_quotes.py — 把券商的 CBAS 選擇權報價單（.xlsx）轉成 data/cbas_quotes.json

用法：
    python scripts/import_quotes.py "C:\\path\\元大證選擇權報價單20261019.xlsx"

需要 openpyxl：
    python -m pip install openpyxl

來源格式（元大證券「轉換公司債資產交換選擇權報價表-債券部」）：
    第 1 列  標題
    第 2 列  B 欄 = 報價日期
    第 4 列  表頭
    第 6 列起 資料；表尾有「註：」等說明列，需濾掉

欄位對應（0-based）：
    0 可轉債名稱   1 可轉債代碼   2 擔保銀行/TCRI   3 權利金(佰元報價)
    4 選擇權到期日 5 賣回價格(A)  6 賣回日期(B)     7 剩餘年期
    8 折現率(C)    9 權利金參考價 10 現股參考價     11 轉換價格
    12 CB轉換價值  13 CB參考價    14 溢(折)價%     15 股價波動率120天%
    16 股價波動率240天%  17 流通餘額占發行總額%  18 原始發行量  19 產業

注意：
  * 債券代號多為 5 碼，但第 10 次以後發行者為 6 碼（例：岳豐十 = 622010），
    過濾時務必接受 5~6 碼，否則會漏抓。
  * 折現率(C) 是券商訂價用的資金成本（如 0.0375 = 3.75%），
    與「持有至賣回的年化報酬(YTP)」是不同的東西，不可混用。
  * 依報價單註九：實際權利金(一張) = (權利金佰元報價 + (CB買進成交價 − 100)) × 1000
"""
import sys, io, os, json, datetime

try:
    import openpyxl
except ImportError:
    sys.exit("需要 openpyxl：python -m pip install openpyxl")

SHEET = "報價單"


def as_num(x):
    return float(x) if isinstance(x, (int, float)) else None


def as_date(x):
    return x.strftime("%Y-%m-%d") if isinstance(x, datetime.datetime) else ""


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    repo = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    dst = os.path.join(repo, "data", "cbas_quotes.json")

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit("找不到工作表「%s」。此檔的工作表：%s" % (SHEET, wb.sheetnames))
    rows = list(wb[SHEET].iter_rows(values_only=True))

    qdate = as_date(rows[1][1]) if len(rows) > 1 else ""
    if not qdate:
        sys.exit("讀不到報價日期（預期在第 2 列 B 欄）")

    quotes, skipped = {}, 0
    for r in rows[4:]:
        code = r[1]
        if code is None:
            continue
        code = str(code).strip()
        # 5~6 碼：第 10 次以後發行的債券代號是 6 碼
        if not (code.isdigit() and 5 <= len(code) <= 6):
            skipped += 1
            continue
        prem = as_num(r[3])
        if prem is None:
            skipped += 1
            continue
        quotes[code] = {
            "name":       str(r[0]).strip() if r[0] else "",
            "premium":    round(prem, 4),                                  # 權利金(佰元報價)
            "discRate":   round(as_num(r[8]) * 100, 4) if as_num(r[8]) is not None else None,
            "putPrice":   as_num(r[5]),
            "putDate":    as_date(r[6]),
            "optExpiry":  as_date(r[4]),
            "convPrice":  as_num(r[11]),                                   # 券商的轉換價（交叉核對用）
            "stockRef":   as_num(r[10]),
            "cbRef":      as_num(r[13]),
            "premiumPct": round(as_num(r[14]), 4) if as_num(r[14]) is not None else None,
            "vol120":     as_num(r[15]),      # 股價波動率 120 天 %
            "vol240":     as_num(r[16]),      # 股價波動率 240 天 %
            "outPct":     as_num(r[17]),      # 流通餘額占發行總額 %
            "industry":   str(r[19]).strip() if len(r) > 19 and r[19] else "",
            "guarantor":  str(r[2]).strip() if r[2] else "",
        }

    out = {
        "_說明": "由 scripts/import_quotes.py 從券商報價單產生，請勿手改；重跑腳本即覆寫。",
        "_公式": "實際權利金(一張) = (權利金佰元報價 + (CB買進成交價 − 100)) × 1000",
        "asOf": qdate,
        "source": "元大證券 轉換公司債資產交換選擇權報價表",
        "sourceFile": os.path.basename(src),
        "parFloor": 100,
        "quotes": quotes,
    }
    with io.open(dst, "w", encoding="utf-8") as f:
        f.write(json.dumps(out, ensure_ascii=False, indent=1))

    print("報價日 %s ─ 取得 %d 檔，略過 %d 列（表尾說明與無報價列）" % (qdate, len(quotes), skipped))
    print("寫入 %s" % dst)

    # 與 bonds.json 比對涵蓋率，順便抓出對不上的代號
    bp = os.path.join(repo, "data", "bonds.json")
    if os.path.exists(bp):
        bonds = json.load(io.open(bp, encoding="utf-8"))
        known = set(b["bondCode"] for b in bonds)
        hit = [c for c in quotes if c in known]
        miss = [c for c in quotes if c not in known]
        print("對得上 bonds.json：%d / %d" % (len(hit), len(quotes)))
        if miss:
            print("對不上（可能已下櫃或代號有異）：%s" % ", ".join(
                "%s %s" % (c, quotes[c]["name"]) for c in miss))


if __name__ == "__main__":
    main()
